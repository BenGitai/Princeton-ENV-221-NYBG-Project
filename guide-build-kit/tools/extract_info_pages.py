"""Extract species info pages built in the ENV 221 info-page template (Madison's
slides: title, "INVASIVE: native to ..." line, IMPACT box, native alternatives with
photos) from a .pptx. One slide = one species. Text and images are copied verbatim.

Usage: python extract_info_pages.py <deck.pptx> <out.json> <photo_dir>
"""
import json, os, re, sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

def slug(s): return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")

def parse_slide(slide, photo_dir):
    texts, pics = [], []
    for sh in slide.shapes:
        if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
            pics.append(sh)
        elif sh.has_text_frame and sh.text_frame.text.strip():
            texts.append(sh)
    page = {"source": "ENV 221 info page (pptx)", "alternatives": []}
    labels = []
    for sh in texts:
        t = sh.text_frame.text.strip()
        if t.upper().startswith("INVASIVE"):
            page["origin_line"] = t
            m = re.search(r"native to (.+)$", t, re.I)
            page["native_range"] = m.group(1).strip() if m else ""
        elif t.upper().startswith("IMPACT"):
            body = t[len("IMPACT"):].strip()
            src = re.search(r"\(Source:\s*(.+?)\)\s*$", body, re.S)
            page["impact"] = re.sub(r"\(Source:.+\)\s*$", "", body, flags=re.S).strip()
            page["impact_source"] = src.group(1).strip() if src else ""
        elif t.lower().startswith("native alternatives"):
            continue
        elif t.lower().startswith("photo:"):
            page["photo_credit"] = t
        elif "(" in t and sh.top < 900000 and "title" not in page:
            m = re.match(r"(.+?)\s*\((.+)\)", t)
            page["title"] = t
            page["scientific_name"], page["common_names"] = m.group(1).strip(), [c.strip() for c in m.group(2).split(",")]
        else:
            labels.append(sh)
    if not page.get("scientific_name"):
        return None
    sid = slug(page["scientific_name"])
    os.makedirs(photo_dir, exist_ok=True)
    def save(pic, name):
        ext = pic.image.ext
        fn = f"{name}.{ext}"
        open(os.path.join(photo_dir, fn), "wb").write(pic.image.blob)
        return fn
    # main photo = the tallest picture; alternatives = the rest, matched to the nearest label below them
    pics.sort(key=lambda p: -p.height)
    if pics:
        page["photo"] = save(pics[0], sid)
    for lab in labels:
        name = lab.text_frame.text.strip()
        near = min(pics[1:], key=lambda p: abs((p.left + p.width / 2) - (lab.left + lab.width / 2)) + abs(lab.top - (p.top + p.height)), default=None)
        page["alternatives"].append({"name": name, "photo": save(near, f"{sid}--{slug(name)}") if near else None})
    page["id"] = sid
    return page

if __name__ == "__main__":
    deck, out, photo_dir = sys.argv[1:4]
    pages = [p for p in (parse_slide(s, photo_dir) for s in Presentation(deck).slides) if p]
    json.dump(pages, open(out, "w"), indent=2, ensure_ascii=False)
    print(len(pages), "page(s):", [p["scientific_name"] for p in pages])
